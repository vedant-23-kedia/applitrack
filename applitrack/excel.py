"""Excel and CSV export, for when you would rather not involve Google."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path

from .models import COLUMNS, Application, Status

FILLS = {
    "Offer": "C6EFCE",
    "Interview": "CCE5F6",
    "Assessment": "FFE8BF",
    "Action Required": "FFE5B4",
    "Acknowledged": "F2F2F2",
    "Applied": "F2F2F2",
    "Rejected": "FFC7CE",
    "Ghosted": "E6E6E6",
    "Withdrawn": "E6E6E6",
}


def write_xlsx(apps: list[Application], path: str, summary: dict | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="33404F")

    ws.append(COLUMNS + ["My Notes"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for app in apps:
        row = app.to_row()
        ws.append([row[c] for c in COLUMNS] + [""])
        color = FILLS.get(app.status.value)
        if color:
            ws.cell(row=ws.max_row, column=COLUMNS.index("Status") + 1).fill = PatternFill(
                "solid", fgColor=color
            )

    # Width by longest cell, capped so the Thread URLs do not dominate.
    for idx, name in enumerate(COLUMNS + ["My Notes"], start=1):
        longest = max(
            [len(name)] + [len(str(app.to_row().get(name, ""))) for app in apps] or [len(name)]
        )
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 48)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if summary is not None:
        s = wb.create_sheet("Summary")
        s["A1"] = "Job Application Tracker"
        s["A1"].font = Font(bold=True, size=14)
        s["A2"], s["B2"] = "Last synced", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        s["A4"], s["B4"] = "Total applications", len(apps)
        s["A5"], s["B5"] = "Still in progress", summary.get("active", 0)
        s["A6"], s["B6"] = "Response rate", f"{summary.get('response_rate', 0):.0%}"
        s["A8"], s["B8"] = "Status", "Count"
        s["A8"].font = s["B8"].font = Font(bold=True)
        row = 9
        for status in Status:
            count = summary.get("by_status", {}).get(status.value)
            if count:
                s[f"A{row}"], s[f"B{row}"] = status.value, count
                row += 1
        s.column_dimensions["A"].width = 24
        s.column_dimensions["B"].width = 22

    wb.save(out)
    return str(out.resolve())


def write_csv(apps: list[Application], path: str) -> str:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        for app in apps:
            writer.writerow(app.to_row())
    return str(out.resolve())
